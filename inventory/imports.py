import hashlib
import re
from collections import Counter
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .models import ImportFile, ImportRow
from .services import require

LAYOUTS = {
    "INVENTARIO ALMACEN": {"date": 5, "headers": {1: "NOMBRE DEL PRODUCTO", 2: "LOTE DEL PRODUCTO", 3: "CANTIDAD", 4: "PIEZAS"}},
    "INVENTARIO REACTIVOS": {"date": 4, "headers": {1: "NOMBRE DEL PRODUCTO", 2: "LOTE", 3: "CAJAS"}},
}


def suggest_date(value, epoch):
    if value is None or str(value).strip().upper() in {"", "S/F"}:
        return None, "Caducidad ausente: verifique etiqueta o documente que no aplica."
    if isinstance(value, datetime):
        value = value.date()
    elif isinstance(value, (int, float)):
        try:
            value = from_excel(value, epoch).date()
        except (ValueError, OverflowError, AttributeError):
            return None, "Fecha numerica invalida."
    elif isinstance(value, str):
        match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", value.strip())
        if not match:
            return None, "Fecha de texto incompleta o mal formada; no se corrige automaticamente."
        first, second, year = (int(part) for part in match.groups())
        if first <= 12 and second <= 12 and first != second:
            return None, "Fecha ambigua entre dia/mes y mes/dia. Confirme en la etiqueta."
        month, day = (first, second) if second > 12 or first == second else (second, first)
        try:
            value = date(year, month, day)
        except ValueError:
            return None, "Fecha de texto invalida."
    else:
        return None, "Tipo de fecha no reconocido."
    if not 1990 <= value.year <= 2100:
        return None, "Fecha fuera del rango razonable (1990-2100); verifique la etiqueta."
    return value, ""


def serializable(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


@transaction.atomic
def stage_workbook(*, user, content, filename):
    require(user, "import_inventory")
    if len(content) > 5 * 1024 * 1024:
        raise ValidationError("El archivo no debe superar 5 MB.")
    digest = hashlib.sha256(content).hexdigest()
    existing = ImportFile.objects.filter(digest=digest).first()
    if existing:
        return existing, False
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 25 * 1024 * 1024 or len(archive.infolist()) > 300:
                raise ValidationError("El libro excede el limite de contenido descomprimido.")
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), data_only=True, read_only=False, keep_links=False)
    except (BadZipFile, KeyError, ValueError, OSError) as error:
        raise ValidationError("No se pudo abrir el archivo XLSX.") from error
    try:
        sheet_names = [name for name in LAYOUTS if name in workbook.sheetnames]
        if not sheet_names:
            raise ValidationError("No se encontraron las hojas INVENTARIO ALMACEN o INVENTARIO REACTIVOS.")
        entries = []
        for name in sheet_names:
            sheet = workbook[name]
            layout = LAYOUTS[name]
            if sheet.max_row > 2001 or sheet.max_column > 50:
                raise ValidationError("La hoja supera el limite de 2000 filas o 50 columnas.")
            if any(str(sheet.cell(1, column).value).strip().upper() != title for column, title in layout["headers"].items()):
                raise ValidationError(f"Los encabezados de {name} no coinciden con el formato revisado.")
            for row_number in range(2, sheet.max_row + 1):
                product = str(sheet.cell(row_number, 1).value or "").strip()
                if not product:
                    continue
                lot_cell = sheet.cell(row_number, 2)
                expiry_cell = sheet.cell(row_number, layout["date"])
                suggested, date_issue = suggest_date(expiry_cell.value, workbook.epoch)
                issues = ["Confirme saldo fisico y equivalencia de cajas/piezas; no se ejecutan las formulas del Excel."]
                if date_issue:
                    issues.append(date_issue)
                if lot_cell.is_date:
                    issues.append("El lote tiene formato de fecha: confirme el identificador original en la etiqueta.")
                if not lot_cell.value or str(lot_cell.value).strip().upper() == "S/L":
                    issues.append("Lote ausente; verifique etiqueta o asigne identificador interno documentado.")
                quantity_columns = [3, 4] if name == "INVENTARIO ALMACEN" else [3]
                quantities = {sheet.cell(1, column).value: {
                    "raw": serializable(sheet.cell(row_number, column).value),
                    "cached": serializable(cached[name].cell(row_number, column).value),
                } for column in quantity_columns}
                original = {"product": product, "lot": str(serializable(lot_cell.value) or ""),
                            "lot_is_date": lot_cell.is_date, "expiry": serializable(expiry_cell.value), "quantities": quantities}
                entries.append({"sheet": name, "row_number": row_number, "original": original,
                                "suggested_expiry": suggested, "issues": issues})
        if not entries:
            raise ValidationError("El libro no contiene filas con productos.")
        counts = Counter((entry["original"]["product"].casefold(), entry["original"]["lot"]) for entry in entries)
        imported = ImportFile.objects.create(name=filename[:200], digest=digest, actor=user)
        for entry in entries:
            if counts[(entry["original"]["product"].casefold(), entry["original"]["lot"])] > 1:
                entry["issues"].append("Producto/lote repetido en el libro. Concilie filas y caducidades antes de importar.")
            ImportRow.objects.create(file=imported, **entry)
        return imported, True
    finally:
        workbook.close()
        cached.close()